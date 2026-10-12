"""Reading text out of source documents, with page and section positions (M2-02)."""

import io
from datetime import UTC, datetime

import docx
import pytest

from apps.reports.pdf import render_report
from apps.sources import parsing

MARKDOWN = """Intro line before any heading.

# Checkout

A shopper can pay by card.

## Refunds

A refund must go back to the card that paid.
"""


def test_markdown_headings_start_sections_and_positions_point_into_the_text():
    parsed = parsing.parse_text(MARKDOWN)
    assert [s["heading"] for s in parsed.sections] == ["", "Checkout", "Refunds"]
    refunds = parsed.sections[2]
    assert parsed.text[refunds["start"] : refunds["end"]].startswith("## Refunds")
    assert parsed.pages == [{"page": 1, "start": 0, "end": len(parsed.text)}]
    # Sections cover the text with no gap and no overlap.
    assert parsed.sections[0]["start"] == 0 and parsed.sections[-1]["end"] == len(parsed.text)
    assert all(
        a["end"] == b["start"] for a, b in zip(parsed.sections, parsed.sections[1:], strict=False)
    )


def test_plain_text_is_one_section_and_windows_line_ends_are_normalised():
    parsed = parsing.parse_text("A shopper can pay by card.\r\nA refund goes back.\r\n")
    assert parsed.text == "A shopper can pay by card.\nA refund goes back."
    assert len(parsed.sections) == 1


def test_a_long_section_is_cut_at_a_paragraph():
    paragraph = "A shopper can pay by card. " * 40  # about 1,000 characters
    parsed = parsing.parse_text("\n\n".join([paragraph] * 30))
    assert len(parsed.sections) > 1
    assert all(s["end"] - s["start"] <= parsing.SECTION_CHARS for s in parsed.sections)
    assert all(
        parsed.text[s["start"] : s["end"]].strip().endswith("card.") for s in parsed.sections
    )


def test_empty_text_is_unreadable():
    with pytest.raises(parsing.Unreadable):
        parsing.parse_text("  \n ")


def _docx(*paragraphs) -> io.BytesIO:
    document = docx.Document()
    for style, text in paragraphs:
        document.add_heading(text, level=1) if style == "h" else document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    out.seek(0)
    return out


def test_a_docx_keeps_its_headings_as_sections():
    parsed = parsing.parse_file(
        _docx(
            ("h", "Checkout"), ("p", "A shopper can pay by card."), ("h", "Refunds"), ("p", "Back.")
        ),
        "docx",
    )
    assert [s["heading"] for s in parsed.sections] == ["Checkout", "Refunds"]
    assert "A shopper can pay by card." in parsed.text


def test_a_pdf_is_read_page_by_page():
    pdf = render_report(
        project="Checkout",
        release="2.4.1",
        report_id="RPT-001",
        body={"summary": "Card payments must be confirmed.", "sections": []},
        signed_by="Ngozi Eze",
        signed_at=datetime(2026, 10, 12, tzinfo=UTC),
        readiness="ready",
        reason="",
    )
    parsed = parsing.parse_file(io.BytesIO(pdf), "pdf")
    assert "Card payments must be confirmed." in " ".join(parsed.text.split())
    assert parsed.pages[0]["page"] == 1 and parsed.sections[0]["heading"] == "Page 1"
    page = parsed.pages[0]
    assert parsed.text[page["start"] : page["end"]].strip()


def test_files_that_cannot_be_read_say_why():
    with pytest.raises(parsing.Unreadable) as bad_pdf:
        parsing.parse_file(io.BytesIO(b"%PDF-1.4 not really a pdf"), "pdf")
    assert "could not be read" in bad_pdf.value.message
    with pytest.raises(parsing.Unreadable) as not_utf8:
        parsing.parse_file(io.BytesIO(b"\xff\xfe\x00caf\xe9"), "text")
    assert "UTF-8" in not_utf8.value.message
    with pytest.raises(parsing.Unreadable):
        parsing.parse_file(io.BytesIO(b""), "png")
