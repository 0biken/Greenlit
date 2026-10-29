"""Extract approved PRD uploads into source documents (M2-02)."""

from pathlib import PurePath
from typing import TypedDict

import pdfplumber
from django.db import transaction
from docx import Document

from apps.core import audit
from apps.core.jobs import JobFailure, job_handler
from apps.files.models import FileState, StoredFile
from apps.files.storage import get_storage
from apps.requirements.models import SourceDocument

MAX_EXTRACTED_CHARS = 10_000_000


class PageIndex(TypedDict):
    page: int | None
    start_char: int
    end_char: int


def _append_page(pages: list[PageIndex], text: str, page: int | None) -> None:
    start = pages[-1]["end_char"] + 1 if pages else 0
    pages.append({"page": page, "start_char": start, "end_char": start + len(text)})


def _pdf(file) -> tuple[str, list[PageIndex], int]:
    pages: list[PageIndex] = []
    text_pages: list[str] = []
    with pdfplumber.open(file) as document:
        for number, page in enumerate(document.pages, start=1):
            text = page.extract_text() or ""
            _append_page(pages, text, number)
            text_pages.append(text)
    return "\n".join(text_pages), pages, len(pages)


def _docx(file) -> tuple[str, list[PageIndex], None]:
    document = Document(file)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    paragraphs.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if cell.text
    )
    text = "\n".join(paragraphs)
    index: list[PageIndex] = [{"page": None, "start_char": 0, "end_char": len(text)}]
    return text, index, None


def _plain_text(file) -> tuple[str, list[PageIndex], None]:
    text = file.read().decode("utf-8-sig")
    index: list[PageIndex] = [{"page": None, "start_char": 0, "end_char": len(text)}]
    return text, index, None


@job_handler("extract_prd", max_retries=3)
def extract_prd(job, file_id: str) -> None:
    stored = StoredFile.objects.filter(pk=file_id).first()
    if stored is None or stored.state != FileState.AVAILABLE or stored.kind != "prd":
        return
    if SourceDocument.objects.filter(stored_file=stored).exists():
        return

    job.set_progress(1, 2, "Extracting source text")
    page_count: int | None
    with get_storage().open(stored.storage_key) as file:
        try:
            if stored.detected_type == "pdf":
                text, index, page_count = _pdf(file)
            elif stored.detected_type == "docx":
                text, index, page_count = _docx(file)
            elif stored.detected_type in {"markdown", "text"}:
                text, index, page_count = _plain_text(file)
            else:
                raise JobFailure("This PRD file type cannot be extracted.")
        except (UnicodeDecodeError, ValueError) as exc:
            raise JobFailure("The uploaded PRD could not be read.") from exc

    if len(text) > MAX_EXTRACTED_CHARS:
        raise JobFailure("The extracted PRD is too large to process.")
    if not text.strip():
        raise JobFailure("No selectable text was found in this PRD.")

    with transaction.atomic():
        source = SourceDocument.objects.create(
            org_id=stored.org_id,
            project=stored.project,
            kind=SourceDocument.Kind.PRD,
            title=PurePath(stored.original_name).stem[:200] or "PRD",
            content_text=text,
            page_count=page_count,
            extraction_index=index,
            stored_file=stored,
            created_by=stored.created_by,
        )
        audit.record(
            "source_document.extracted",
            "source_document",
            source.pk,
            actor=stored.created_by,
            after={"project_id": str(stored.project_id), "file_id": str(stored.pk)},
        )
    job.set_progress(2, 2, "Source ready")
